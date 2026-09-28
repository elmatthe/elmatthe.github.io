"""Opt-in real Excel regression; ordinary pytest never launches Excel."""
from copy import copy
from datetime import date
import hashlib
import os
from pathlib import Path
import shutil
import warnings
import zipfile
import io
from xml.etree import ElementTree as ET
import openpyxl
from openpyxl.worksheet.cell_range import CellRange
import pytest
from cpi_config import ROOT, load_config
from workbook_io import write_to_workbook

CANONICAL_SHA256 = '7df41a06452c854b3a4f5536b7b648b64d936513a999a3c2eaa12cddc250f8c1'

def test_reference_schema_matches_default():
    config=load_config(ROOT/'scripts/config-default.toml')
    wb=openpyxl.load_workbook(ROOT/'files/CPI.xlsx',read_only=True,data_only=False)
    try:
        assert tuple(next(wb['CPI_Downloader'].iter_rows(min_row=1,max_row=1,max_col=19,values_only=True)))==config.headers
        cells={openpyxl.utils.get_column_letter(i):getattr(c.value,'text',c.value) for i,c in enumerate(next(wb['US Data'].iter_rows(min_row=5,max_row=5)),1)}
        for output,raw in [('AI','O'),('AJ','P'),('AK','Q')]:
            assert f'CPI_Downloader!${raw}$2:${raw}$5000' in cells[output]
        for output in ['AM','AN','AO','AP','AQ','AR','AS']:
            assert 'CPI_Downloader!$R$2:$R$5000' in cells[output]
        for output in ['AU','AV','AW','AX','AY','AZ','BA']:
            assert 'CPI_Downloader!$S$2:$S$5000' in cells[output]
        assert 'CPI_Downloader!$B$2:$S$5000' in cells['A']
    finally:
        wb.close()

def formula_value(value):
    # Modern array/dynamic formulas are objects, not plain strings.
    if hasattr(value, '__dict__'):
        return (type(value).__name__, dict(vars(value)))
    return value

def snapshot(path):
    with warnings.catch_warnings():
        warnings.simplefilter('ignore', UserWarning)
        wb = openpyxl.load_workbook(path, data_only=False)
        cached = openpyxl.load_workbook(path, data_only=True)
    sheets = {}
    for ws in wb:
        cells = {}
        styles = {}
        array_ranges = [CellRange(ref) for ref in ws.array_formulae.values()]
        for cell in list(ws._cells.values()):
            if cell.value is not None or cell.has_style:
                if cell.style_id not in styles:
                    styles[cell.style_id] = (copy(cell.font), copy(cell.fill), copy(cell.border), copy(cell.alignment), cell.number_format, copy(cell.protection))
                value = cell.value
                # Excel can serialize array interior cells with an empty <f/>.
                # openpyxl exposes those markers as '='; saving in Excel removes
                # them. Compare their actual cached values, only inside a known
                # array range; retain the real anchor formula and all formatting.
                if value == '=':
                    assert any(cell.coordinate in r and (cell.row,cell.column)!=(r.min_row,r.min_col) for r in array_ranges), 'Empty formula outside array interior'
                    value = cached[ws.title][cell.coordinate].value
                cells[cell.coordinate] = (formula_value(value), *styles[cell.style_id])
        sheets[ws.title] = {
            'cells':cells,
            'cached_values':{c.coordinate:c.value for c in cached[ws.title]._cells.values() if c.value is not None},
            'rows': {k:dict(v) for k,v in ws.row_dimensions.items()},
            'cols': {k:dict(v) for k,v in ws.column_dimensions.items()},
            'merges':str(ws.merged_cells),
            'tables':[openpyxl.xml.functions.tostring(t.to_tree()) for t in ws.tables.values()],
            'conditional':[ (str(k), [openpyxl.xml.functions.tostring(r.to_tree()) for r in rules]) for k,rules in ws.conditional_formatting._cf_rules.items()],
            'charts':[openpyxl.xml.functions.tostring(c._write()) for c in ws._charts],
            'state':ws.sheet_state,
        }
    result = {'sheets':sheets,'names':[(k,dict(v)) for k,v in wb.defined_names.items()], 'calculation':dict(wb.calculation)}
    wb.close()
    cached.close()
    return result

def test_array_interior_serialization_preserves_real_values(tmp_path):
    path=tmp_path/'array.xlsx'
    wb=openpyxl.Workbook()
    ws=wb.active
    ws['A1']=openpyxl.worksheet.formula.ArrayFormula(ref='A1:A2',text='=SEQUENCE(2)')
    ws['A2']=2
    wb.save(path);wb.close()
    expected=snapshot(path)
    original=path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path,'w') as target:
        for entry in source.infolist():
            data=source.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(data)
                cell=root.find(".//{*}c[@r='A2']")
                cell.insert(0,ET.Element('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f',ca='1'))
                data=ET.tostring(root)
            target.writestr(entry,data)
    assert snapshot(path)==expected
    # A real value or anchor-formula change must still fail preservation.
    wb=openpyxl.load_workbook(io.BytesIO(original))
    wb.active['A2']=3;wb.save(path)
    assert snapshot(path)!=expected
    wb.active['A2']=2
    wb.active['A1']=openpyxl.worksheet.formula.ArrayFormula(ref='A1:A2',text='=SEQUENCE(2)+1')
    wb.save(path);wb.close()
    assert snapshot(path)!=expected

def test_empty_formula_outside_array_rejected(tmp_path):
    path=tmp_path/'invalid.xlsx'
    wb=openpyxl.Workbook();wb.active['A1']='=';wb.save(path);wb.close()
    with pytest.raises(AssertionError,match='Empty formula outside array'):
        snapshot(path)

@pytest.mark.skipif(os.environ.get('CPI_EXCEL_REGRESSION') != '1', reason='Opt-in desktop Excel test')
def test_canonical_copy_preserves_model(tmp_path):
    canonical=ROOT/'files/CPI.xlsx'
    digest=lambda:hashlib.sha256(canonical.read_bytes()).hexdigest()
    initial=digest()
    assert initial==CANONICAL_SHA256
    target=tmp_path/'CPI-regression.xlsx'
    shutil.copy2(canonical,target)
    before=snapshot(target)
    config=load_config(ROOT/'scripts/config-default.toml')
    rows=[{'date':date(2024,m,1), **{c.id:float(i+m) for i,c in enumerate(config.columns) if c.id!='date'}} for m in (1,2)]
    try:
        write_to_workbook(rows,target,config)
        after=snapshot(target)
        assert before['names']==after['names']
        assert before['calculation']==after['calculation']
        assert before['sheets'].keys()==after['sheets'].keys()
        for name, old in before['sheets'].items():
            new=after['sheets'][name]
            if name!='CPI_Downloader':
                assert old==new, f'Model formulas, values, formatting or structure changed: {name}'
                continue
            assert {k:v for k,v in old.items() if k not in ('cells','cached_values')}=={k:v for k,v in new.items() if k not in ('cells','cached_values')}
            for coord, cell in old['cells'].items():
                assert cell[1:]==new['cells'].get(coord,(None,)+cell[1:])[1:], f'Raw formatting changed: {coord}'
        wb=openpyxl.load_workbook(target,data_only=False)
        ws=wb[config.sheet_name]
        assert tuple(ws.cell(1,c).value for c in range(1,20))==config.headers
        for index,row in enumerate(rows,2):
            for col,value in enumerate(config.values(row),1):
                actual=ws.cell(index,col).value
                if isinstance(value,date):
                    actual=actual.date() if hasattr(actual,'date') else actual
                assert actual==value
        assert all(ws.cell(r,c).value is None for r in range(4,ws.max_row+1) for c in range(1,20))
        assert tuple(c.id for c in config.columns[15:])==('us_treasury_10y_avg','us_treasury_2y_avg','us_pce_sa','us_core_pce_sa')
        wb.close()
    finally:
        assert digest()==initial, 'Canonical workbook changed'
